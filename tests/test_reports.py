"""reports/: bugün özeti, haftalık özet, sol panel ajandası ve CSV/Excel çıktısı."""

from datetime import date, timedelta

from isyonetim.models import ecommerce as eco
from isyonetim.models import meeting as mtg
from isyonetim.models import project as proj
from isyonetim.reports import export
from isyonetim.reports.today import build_agenda_lists, build_today_summary, format_today_digest
from isyonetim.reports.weekly import build_weekly_summary

TODAY = date(2026, 6, 15)


def d(delta: int) -> str:
    return (TODAY + timedelta(days=delta)).isoformat()


# -------------------------------------------------------------- bugün özeti

def test_today_summary_collects_all_three_modules(conn):
    p = proj.create_project(conn, "P")
    proj.create_task(conn, p.id, "Öncelikli", priority_bucket="today")
    proj.create_task(conn, p.id, "Vadesi geçmiş", due_date=d(-2))
    proj.create_task(conn, p.id, "İleride", due_date=d(5))
    mtg.create_meeting(conn, "Bugünkü", d(0))
    mtg.create_meeting(conn, "Yarınki", d(1))
    m = mtg.create_meeting(conn, "Kaynak", d(-10))
    mtg.add_action_item(conn, m.id, "Geciken aksiyon", owner="A", due_date=d(-1))
    mtg.add_action_item(conn, m.id, "Bugün vadeli", owner="A", due_date=d(0))

    s = build_today_summary(conn, today=TODAY)

    assert s["date"] == d(0)
    assert {t.title for t in s["tasks"]} == {"Öncelikli", "Vadesi geçmiş"}
    assert [m.title for m in s["meetings_today"]] == ["Bugünkü"]
    assert [a.description for a in s["overdue_actions"]] == ["Geciken aksiyon"]
    assert [a.description for a in s["due_today_actions"]] == ["Bugün vadeli"]


def test_today_summary_does_not_duplicate_task_in_both_buckets(conn):
    """'today' önceliğinde VE vadesi geçmiş bir görev iki kez listelenmemeli."""
    p = proj.create_project(conn, "P")
    proj.create_task(conn, p.id, "İkisi birden", priority_bucket="today", due_date=d(-1))

    s = build_today_summary(conn, today=TODAY)

    assert len(s["tasks"]) == 1


def test_completed_tasks_excluded_from_today(conn):
    p = proj.create_project(conn, "P")
    t = proj.create_task(conn, p.id, "Bitti", priority_bucket="today")
    proj.update_task_status(conn, t.id, "completed")

    assert build_today_summary(conn, today=TODAY)["tasks"] == []


def test_digest_text_for_empty_and_full_day(conn):
    empty = build_today_summary(conn, today=TODAY)
    assert "yok" in format_today_digest(empty)

    p = proj.create_project(conn, "P")
    proj.create_task(conn, p.id, "Görev", priority_bucket="today")
    mtg.create_meeting(conn, "Toplantı", d(0))

    text = format_today_digest(build_today_summary(conn, today=TODAY))
    assert "1 görev" in text
    assert "1 toplantı" in text


# -------------------------------------------------------------- haftalık özet

def test_weekly_revenue_and_profit(conn):
    p = eco.create_product(conn, "Ürün", cost_price=100)
    c = eco.create_channel(conn, "Kanal")
    eco.add_sale(conn, p.id, c.id, d(-1), quantity=2, unit_price=150)
    eco.add_sale(conn, p.id, c.id, d(-3), quantity=1, unit_price=150, unit_cost=120)

    s = build_weekly_summary(conn, start=TODAY - timedelta(days=6), end=TODAY)

    assert s["sales_count"] == 2
    assert s["total_revenue"] == 450          # 2*150 + 1*150
    assert s["total_profit"] == 130           # (150-100)*2 + (150-120)*1
    assert s["channel_totals"][0].channel_name == "Kanal"
    assert s["channel_totals"][0].quantity == 3


def test_weekly_profit_unknown_when_any_cost_missing(conn):
    """Maliyeti bilinmeyen tek bir satış bile varsa toplam kâr 'bilinmiyor'
    olmalı; eksik veriyle yanlış rakam göstermekten iyidir."""
    p = eco.create_product(conn, "Maliyetsiz")      # cost_price yok
    c = eco.create_channel(conn, "Kanal")
    eco.add_sale(conn, p.id, c.id, d(-1), quantity=1, unit_price=150)

    s = build_weekly_summary(conn, start=TODAY - timedelta(days=6), end=TODAY)

    assert s["total_revenue"] == 150
    assert s["total_profit"] is None
    assert s["top_selling_products"][0].profit is None


def test_weekly_sales_outside_period_excluded(conn):
    p = eco.create_product(conn, "Ürün", cost_price=10)
    c = eco.create_channel(conn, "Kanal")
    eco.add_sale(conn, p.id, c.id, d(-30), quantity=5, unit_price=100)

    s = build_weekly_summary(conn, start=TODAY - timedelta(days=6), end=TODAY)

    assert s["sales_count"] == 0
    assert s["total_revenue"] == 0


def test_weekly_task_buckets(conn):
    p = proj.create_project(conn, "P")
    done = proj.create_task(conn, p.id, "Tamamlanan")
    proj.update_task_status(conn, done.id, "completed")
    running = proj.create_task(conn, p.id, "Devam")
    proj.update_task_status(conn, running.id, "in_progress")
    proj.create_task(conn, p.id, "Geciken", due_date="2020-01-01")

    s = build_weekly_summary(conn)

    assert [t.title for t in s["in_progress_tasks"]] == ["Devam"]
    assert [t.title for t in s["overdue_tasks"]] == ["Geciken"]
    # completed_at bugün damgalandı; varsayılan dönem son 7 gün
    assert [t.title for t in s["completed_tasks"]] == ["Tamamlanan"]


def test_top_products_limited_to_five(conn):
    c = eco.create_channel(conn, "Kanal")
    for i in range(7):
        p = eco.create_product(conn, f"Ürün {i}", cost_price=1)
        eco.add_sale(conn, p.id, c.id, d(-1), quantity=i + 1, unit_price=10)

    s = build_weekly_summary(conn, start=TODAY - timedelta(days=6), end=TODAY)

    assert len(s["top_selling_products"]) == 5
    assert s["top_selling_products"][0].quantity == 7      # en çok satan başta


# ------------------------------------------------------------ ajanda (sol panel)

def test_agenda_splits_upcoming_and_overdue(conn):
    today = date.today()
    p = proj.create_project(conn, "P")
    proj.create_task(conn, p.id, "Yaklaşan görev", due_date=(today + timedelta(days=2)).isoformat())
    proj.create_task(conn, p.id, "Geciken görev", due_date=(today - timedelta(days=2)).isoformat())
    proj.create_task(conn, p.id, "Çok ileride", due_date=(today + timedelta(days=60)).isoformat())
    mtg.create_meeting(conn, "Yaklaşan toplantı", (today + timedelta(days=1)).isoformat(), meeting_time="10:00")

    agenda = build_agenda_lists(conn)

    upcoming = " ".join(i["etiket"] for i in agenda["upcoming"])
    overdue = " ".join(i["etiket"] for i in agenda["overdue"])
    assert "Yaklaşan görev" in upcoming
    assert "Yaklaşan toplantı" in upcoming
    assert "Çok ileride" not in upcoming
    assert "Geciken görev" in overdue


def test_agenda_respects_limit(conn):
    today = date.today()
    p = proj.create_project(conn, "P")
    for i in range(10):
        proj.create_task(conn, p.id, f"Görev {i}", due_date=(today + timedelta(days=1)).isoformat())

    assert len(build_agenda_lists(conn, limit=3)["upcoming"]) == 3


def test_completed_tasks_absent_from_agenda(conn):
    today = date.today()
    p = proj.create_project(conn, "P")
    t = proj.create_task(conn, p.id, "Bitti", due_date=(today + timedelta(days=1)).isoformat())
    proj.update_task_status(conn, t.id, "completed")

    assert build_agenda_lists(conn)["upcoming"] == []


# ------------------------------------------------------------------ dışa aktarım

def test_csv_export_writes_header_and_rows(conn, tmp_path):
    p = proj.create_project(conn, "P")
    proj.create_task(conn, p.id, "Görev başlığı")

    out = export.export_rows(proj.list_tasks(conn), tmp_path / "gorevler.csv")

    text = out.read_text(encoding="utf-8-sig")
    assert "title" in text.splitlines()[0]
    assert "Görev başlığı" in text


def test_csv_export_handles_empty_list(tmp_path):
    out = export.export_rows([], tmp_path / "bos.csv")
    assert out.exists()


def test_excel_export_produces_valid_workbook(conn):
    import io

    import openpyxl

    from isyonetim.reports import excel_export

    p = proj.create_project(conn, "P")
    proj.create_task(conn, p.id, "Görev", priority_bucket="today")

    data = excel_export.build_today_excel(build_today_summary(conn, today=date.today()))
    wb = openpyxl.load_workbook(io.BytesIO(data))

    assert "Görevler" in wb.sheetnames
    assert "Toplantılar" in wb.sheetnames


def test_weekly_excel_has_all_sheets(conn):
    import io

    import openpyxl

    from isyonetim.reports import excel_export

    data = excel_export.build_weekly_excel(build_weekly_summary(conn))
    wb = openpyxl.load_workbook(io.BytesIO(data))

    assert "Kanal Bazlı Satış" in wb.sheetnames
    assert "En Çok Kâr Eden" in wb.sheetnames
